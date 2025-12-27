# Copyright (c) 2025, TNTSWA and contributors
# For license information, please see license.txt

import frappe
from frappe import _


def execute(filters=None):
	columns = get_columns()
	data = get_data(filters)
	return columns, data

def get_columns():
	return [
		_("Name") + ":Data:200",
		_("District") + ":Link/District:200",
		_("Job Description") + ":Data:200",
		_("Amount") + ":Data:200",
	]

def get_data(filters):
	data = []
	item_filters = {
		"docstatus": ["!=", 2]   
	}

	if filters.get("district"):
		item_filters["district"] = filters.get("district")

	members = frappe.get_all(
		"Payment",
		filters=item_filters,
		fields=[
			"member_name",
			"district",
			"amount",
			"membership_id",
			"name_of_non_member"
		]
	)

	for i in members:
		job=frappe.db.get_value("Membership Form",{'name':i.membership_id},'job_description')
		row = [
			i.member_name or i.name_of_non_member,
			i.district,
			job or '',
			i.amount
		]
		data.append(row)

	return data




import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.image import Image
import requests
from io import BytesIO
import openpyxl
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
import openpyxl.styles as styles
from openpyxl.drawing.image import Image
import requests
import PIL
import io
import urllib3


@frappe.whitelist()
def download_member_excel():

	wb = Workbook()
	ws = wb.active
	ws.title = "Payment Members"
	headers = ["வ. எண்", "புகைப்படம்", "", "பெயர்", "பொறுப்பு", "நன்கொடை ரூபாய்"]
	ws.append(headers)

	# Styles
	header_font = Font(bold=True)
	header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
	center_align = Alignment(horizontal="center", vertical="center")

	# Style header
	for col in range(1, len(headers) + 1):
		cell = ws.cell(row=1, column=col)
		cell.font = Font(bold=True, color="FFFFFF")  
		cell.fill = header_fill
		cell.alignment = center_align

	ws.merge_cells("B1:C1")

	ws.column_dimensions["A"].width = 5   
	ws.column_dimensions["B"].width = 7  
	ws.column_dimensions["C"].width = 15 
	ws.column_dimensions["D"].width = 25  
	ws.column_dimensions["E"].width = 20 
	ws.column_dimensions["F"].width = 18 

	thin = Side(border_style="thin", color="000000")
	border = Border(left=thin, right=thin, bottom=thin)

	payments = frappe.get_all(
		"Payment",
		fields=["member_name", "amount", "district", "membership_id", "name_of_non_member"]
	)

	row_no = 2
	serial = 1

	for p in payments:

			ws.cell(row=row_no, column=1).value = serial
			ws.cell(row=row_no, column=6).value = p.amount

			if p.membership_id:
				# MEMBER
				mem = frappe.get_doc("Membership Form", p.membership_id)

				ws.cell(row=row_no, column=4).value = p.member_name
				ws.cell(row=row_no, column=5).value = mem.job_description or ''

				if mem.image:
					image_url = 'https://tntswa.teamproit.com/' + mem.image
					try:
						image_data = requests.get(image_url).content
						img = Image(io.BytesIO(image_data))
						img.width = 80
						img.height = 100

						ws.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=3)
						ws.add_image(img, f'C{row_no}')

						ws.row_dimensions[row_no].height = (img.height * 0.75) + 10

					except Exception as e:
						frappe.log_error(f"Image Error: {e}")

			else:
				# NON-MEMBER
				ws.cell(row=row_no, column=4).value = p.name_of_non_member
				ws.cell(row=row_no, column=5).value = ''   # no job
				# no image

			# borders & alignment
			for col in range(1, 7):
				cell = ws.cell(row=row_no, column=col)
				cell.border = border
				cell.alignment = center_align

			row_no += 1
			serial += 1


	save_path = frappe.get_site_path("public", "files", "payment_members.xlsx")
	wb.save(save_path)

	return "/files/payment_members.xlsx"
