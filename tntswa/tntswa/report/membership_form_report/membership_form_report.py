# Copyright (c) 2025, TNTSWA and contributors
# For license information, please see license.txt

import frappe
from frappe import _




def execute(filters=None):
    columns = get_columns()
    data = get_data(filters)
    return columns, data

def get_columns():
    return [
        _("Name") + ":Data:200",
        _("Shop No") + ":Data:150",
        _("Date of Birth") + ":Date:200",
        _("Age") + ":Data:90",
        _("District") + ":Link/District:200",
        _("Address") + ":Data:450",
    ]

def get_data(filters):
    if not filters:
        filters = {}

    data = []
    item_filters = {
        "docstatus": ["!=", 2]
    }

    if filters.get("district"):
        item_filters["district_name"] = filters.get("district")

    members = frappe.get_all(
        "Membership Form",
        filters=item_filters,
        fields=[
            "name",
            "name1",
            "district_name",
            "shop_no",
            "date_of_birth",
            "age",
            "full_address__permanent_address",
            "image"
        ]
    )

    for i in members:
        dob=frappe.utils.formatdate(i.date_of_birth)
        row = [
            i.name1,
            i.shop_no,
            dob,
            i.age,
            i.district_name,
            i.full_address__permanent_address
        ]
        data.append(row)

    return data


import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from io import BytesIO

@frappe.whitelist()
def download_member_excel(filters=None):
    if filters:
        filters = frappe.parse_json(filters)
    else:
        filters = {}

    columns = ["பெயர்", "கடை எண்", "பிறந்த தேதி", "வயது", "மாவட்டம்", "முகவரி"]
    data = get_data(filters)

    wb = Workbook()
    ws = wb.active
    ws.title = "Members"

    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_num, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num, value=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_num, row_data in enumerate(data, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    frappe.local.response.filename = "payment_members.xlsx"
    frappe.local.response.filecontent = output.read()
    frappe.local.response.type = "download"

    return None   



import frappe
import base64
from frappe.utils.pdf import get_pdf
from frappe.utils import get_site_path, get_url
from datetime import datetime
from frappe.utils import formatdate

@frappe.whitelist()
def download_donation_pdf(district=None):

    conditions = ""
    values = {}

    if district:
        conditions = "AND m.district_name = %(district)s"
        values["district"] = district

    members = frappe.db.sql(f"""
        SELECT
            m.name1,
            m.shop_no,
            m.date_of_birth,
            m.age,
            m.district_name,
            m.full_address__permanent_address
        FROM `tabMembership Form` m
        WHERE m.docstatus != 2
        {conditions}
        ORDER BY m.name1
    """, values, as_dict=True)


    # ---------- BASE64 WATERMARK IMAGE ----------
    logo_path = get_site_path(
        "public", "files", "TNTSWA_LOGO1-removebg-preview.png"
    )
    with open(logo_path, "rb") as f:
        logo_base64 = base64.b64encode(f.read()).decode()

    html = f"""
    <html>
    <head>
    <style>
        @page {{
            size: A4;
            margin: 0;
        }}

        body {{
            font-family: Arial, sans-serif;
            margin: 0;
        }}

        .page {{
            position: relative;
            page-break-after: always;
            width: 210mm;
            height: 247mm;
            overflow: hidden;
        }}

        /* WATERMARK – MOST STABLE METHOD */
        .watermark {{
            position: absolute;
            top: 0;
            left: 0;
            right: 0;
            bottom: 0;

            background-image: url("data:image/png;base64,{logo_base64}");
            background-repeat: no-repeat;
            background-position: center center;
            background-size: 65%;

            z-index: 0;
        }}

        .content {{
            position: relative;
            z-index: 1;
            padding-left: 20mm;
            padding-right:20mm;
            padding-top:5mm;
        }}

        h3 {{
            text-align: center;
            margin: 0 0 5px 0;
        }}

        p {{
            text-align: center;
            font-size: 12px;
            margin: 0;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
        }}

        tr {{
            page-break-inside: avoid;
        }}

        th, td {{
            border: 1px solid #000;
            padding: 6px;
            font-size: 11px;
        }}

        th {{
            background: #FF0000;
            color:white;
        }}

        img {{
            display: block;
            margin: auto;
        }}
    </style>
    </head>
    <body>
    """

    # ---------- PAGINATION ----------
    rows_per_page = 18
    pages = [members[i:i + rows_per_page] for i in range(0, len(members), rows_per_page)]

    sno = 1

    for page in pages:
        html += """
        <div class="page">
            <div class="watermark"></div>
            <div class="content">
            <div>
                <img src="/files/TNTSWA Header.jpeg" alt="TNTSWA Logo" width="700px">
           
        """

        if district:
            html += f"<h4 style=color:white;text-align:center;> {district}</h4></div><br>"
        else: 
            html += f"<h4 style=color:white;text-align:center;> </h4></div>"

        html += """
                <table>
                    <thead>
                        <tr>
                            <th>வ. எண்</th>
                            <th>பெயர்</th>
                            <th>கடை எண்</th>
                            <th>பிறந்த தேதி</th>
                            <th>வயது</th>
                            <th>மாவட்டம்</th>
                            <th>முகவரி</th>
                        </tr>
                    </thead>
                    <tbody>
        """

        for m in page:
            dob = formatdate(m.date_of_birth, "dd-MM-yyyy") if m.date_of_birth else ""
            html += f"""
                <tr>
                    <td align="center">{sno}</td>
                    <td>{m.name1 or ''}</td>
                    <td>{m.shop_no or ''}</td>
                    <td style="font-size:9px;">{dob or ''}</td>
                    <td>{m.age or ''}</td>
                    <td>{m.district_name or ''}</td>
                    <td>{m.full_address__permanent_address or ''}</td>
                </tr>
            """
            sno += 1

        html += """
                    </tbody>
                </table>
            </div>
        </div>
        """

    html += "</body></html>"

    pdf = get_pdf(
        html,
        options={
            "background": True,
            "images": True,
            "print-media-type": True,
            "disable-smart-shrinking": True,
            "margin-top": "0",
            "margin-bottom": "0",
            "margin-left": "0",
            "margin-right": "0",
            "enable-local-file-access": True
        }
    )

    file = frappe.get_doc({
        "doctype": "File",
        "file_name": "Membership_Report.pdf",
        "content": pdf,
        "is_private": 0
    })
    file.save(ignore_permissions=True)

    return get_url(file.file_url)



